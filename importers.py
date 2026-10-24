# -*- coding: utf-8 -*-
"""Извлечение трек-номеров из присланного файла: текст, CSV, Excel (.xlsx)."""

import csv
import io
import re

NUM_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9\-]{5,}$")


class Unsupported(Exception):
    """Формат файла не поддерживается."""


def looks_like_number(value):
    s = str(value).strip()
    return bool(NUM_RE.match(s)) and bool(re.search(r"\d", s))


def _from_text(data):
    text = data.decode("utf-8", "replace")
    out = []
    for token in re.split(r"[\s,;]+", text):
        s = token.strip().strip(".:;!?()[]\"'")
        if s and looks_like_number(s) and s not in out:
            out.append(s)
    return out


def _from_csv(data):
    text = data.decode("utf-8-sig", "replace")
    lines = text.splitlines()
    first = lines[0] if lines else ""
    delim = ";" if first.count(";") >= first.count(",") else ","
    out = []
    for row in csv.reader(io.StringIO(text), delimiter=delim):
        for cell in row:
            s = (cell or "").strip()
            if s and looks_like_number(s) and s not in out:
                out.append(s)
    return out


def _from_xlsx(data):
    try:
        from openpyxl import load_workbook
    except ImportError:  # noqa: BLE001
        raise Unsupported("На сервере не установлен openpyxl — пришлите .txt или .csv.")
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    if isinstance(cell, float) and cell.is_integer():
                        cell = int(cell)
                    s = str(cell).strip()
                    if s and looks_like_number(s) and s not in out:
                        out.append(s)
    finally:
        wb.close()
    return out


def numbers_from_file(filename, data):
    """Возвращает список трек-номеров из файла (по расширению)."""
    ext = (filename or "").lower().rsplit(".", 1)[-1]
    if ext in ("xlsx", "xlsm"):
        return _from_xlsx(data)
    if ext == "xls":
        raise Unsupported("Формат .xls не поддерживается — сохраните как .xlsx или .csv.")
    if ext == "csv":
        return _from_csv(data)
    if ext in ("txt", "text", ""):
        return _from_text(data)
    raise Unsupported("Поддерживаются файлы: .txt, .csv, .xlsx")
